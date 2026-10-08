import json
from types import SimpleNamespace

import pytest
from docx import Document

from app.service import number_check_alignment as alignment
from app.service import number_check_service


def segment(text, index, source="body", row_context=""):
    return SimpleNamespace(text=text, para_index=index, source=source, row_context=row_context)


def response(groups):
    return json.dumps({"pairs": [{"source": s, "target": t} for s, t in groups]})


def test_merged_paragraphs_preserve_numbers_and_target_location(monkeypatch):
    sources = [segment("power: on", 0), segment("button 1", 1), segment("battery 2 AA", 2)]
    targets = [segment("電源：開啟 按鈕 1", 8), segment("電池 3 顆 AA", 9, "table", "電池 3 顆 AA")]
    monkeypatch.setattr(alignment, "generate_text", lambda **_: response([([0, 1], [0]), ([2], [1])]))

    pairs = alignment.align_number_check_segments(sources, targets, model_name="test", log_callback=lambda _: None)

    assert pairs == [
        ("power: on\nbutton 1", "電源：開啟 按鈕 1", 8, "", "body"),
        ("battery 2 AA", "電池 3 顆 AA", 9, "電池 3 顆 AA", "table"),
    ]


def test_split_paragraphs_use_context_instead_of_wrong_single_index(monkeypatch):
    sources = [segment("power: on; button 1", 0)]
    targets = [segment("電源：開啟", 8), segment("按鈕 1", 9, "table")]
    monkeypatch.setattr(alignment, "generate_text", lambda **_: response([([0], [0, 1])]))
    pairs = alignment.align_number_check_segments(sources, targets, model_name="test", log_callback=lambda _: None)
    assert pairs == [("power: on; button 1", "電源：開啟\n按鈕 1", -1, "", "body")]


@pytest.mark.parametrize("groups", [
    [([0], [0])],
    [([0, 0], [0, 1])],
    [([1, 0], [0, 1])],
    [([0, 2], [0, 1])],
    [([False, 1], [0, 1])],
    [([], []), ([0, 1], [0, 1])],
])
def test_invalid_coverage_is_rejected(groups):
    with pytest.raises(ValueError):
        alignment._validate_mapping(response(groups), [0, 1], [0, 1])


@pytest.mark.parametrize("content", ["not json", "```", "[]", '{"pairs": []}'])
def test_malformed_response_is_rejected_with_retryable_error(content):
    with pytest.raises(ValueError):
        alignment._validate_mapping(content, [0], [0])


def test_json_fence_is_accepted():
    content = response([([0], [0])])
    assert alignment._validate_mapping(f"```json\n{content}\n```", [0], [0]) == [
        {"source": [0], "target": [0]},
    ]


def test_retry_then_fail_before_number_check(monkeypatch):
    calls = []

    def invalid_response(**kwargs):
        calls.append(kwargs)
        return response([([0], [0])])

    monkeypatch.setattr(alignment, "generate_text", invalid_response)
    with pytest.raises(ValueError, match="自动生成对照失败"):
        alignment.align_number_check_segments(
            [segment("a", 0), segment("b", 1)], [segment("甲", 0), segment("乙", 1)],
            model_name="selected-model", log_callback=lambda _: None,
        )
    assert len(calls) == 3
    assert calls[0]["model"] == "selected-model"
    assert calls[0]["route"] == "openrouter"


def test_windows_keep_boundary_context_without_omissions(monkeypatch):
    monkeypatch.setattr(alignment, "_WINDOW_SEGMENTS", 3)
    calls = []

    def align_window(**kwargs):
        payload = json.loads(kwargs["user_prompt"])
        ids = [s["id"] for s in payload["source"]]
        calls.append(ids)
        return response([([i], [i]) for i in ids])

    monkeypatch.setattr(alignment, "generate_text", align_window)
    sources = [segment(f"source {i}", i) for i in range(7)]
    targets = [segment(f"译文 {i}", i) for i in range(7)]
    pairs = alignment.align_number_check_segments(sources, targets, model_name="test", log_callback=lambda _: None)

    assert calls == [[0, 1, 2], [2, 3, 4], [4, 5, 6]]
    assert [p[0] for p in pairs] == [s.text for s in sources]
    assert [p[1] for p in pairs] == [t.text for t in targets]


def test_missing_and_added_content_is_preserved(monkeypatch):
    monkeypatch.setattr(alignment, "generate_text", lambda **_: response([
        ([0], [0]), ([1], []), ([], [1]),
    ]))
    pairs = alignment.align_number_check_segments(
        [segment("button 1", 0), segment("battery 2 AA", 1)],
        [segment("按鈕 1", 0), segment("新增 3", 1)],
        model_name="test", log_callback=lambda _: None,
    )
    assert pairs[1] == ("battery 2 AA", "", -1, "", "body")
    assert pairs[2] == ("", "新增 3", 1, "", "body")


@pytest.mark.parametrize("root_name", ["数检_程序-AI", "数检_程序-AIV2"])
def test_docx_merge_reaches_rule_check_and_preserves_table_location(tmp_path, monkeypatch, root_name):
    root = number_check_service.REPO_ROOT / "专检" / root_name
    monkeypatch.setattr(number_check_service.settings, "NUMBER_CHECK_ROOT", str(root))
    monkeypatch.setenv("API_KEY", "test-key")
    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    main = number_check_service._load_latest_main_module()
    source_path, target_path = tmp_path / "source.docx", tmp_path / "target.docx"
    source = Document()
    source.add_paragraph("power: on")
    source.add_paragraph("button 1")
    source.add_table(rows=1, cols=1).cell(0, 0).text = "battery 2 AA"
    source.sections[0].header.paragraphs[0].text = "source header 99"
    source.save(source_path)
    target = Document()
    target.add_paragraph("電源：開啟 按鈕 1")
    target.add_table(rows=1, cols=1).cell(0, 0).text = "電池 3 顆 AA"
    target.sections[0].header.paragraphs[0].text = "译文页眉 98"
    target.save(target_path)
    monkeypatch.setattr(alignment, "generate_text", lambda **_: response([([0, 1], [0]), ([2], [1])]))
    callback = lambda s, t: alignment.align_number_check_segments(
        s, t, model_name="test", log_callback=lambda _: None,
    )

    pairs = main._build_pairs(str(source_path), str(target_path), alignment_callback=callback)
    rows = main.check_text_pairs(pairs)

    assert len(pairs) == 2
    assert pairs[1][2:] == (1, "電池 3 顆 AA", "table")
    assert rows[0]["是否错误"] == "✔正确"
    assert rows[1]["是否错误"] == "❗错误"
    assert all("99" not in pair[0] and "98" not in pair[1] for pair in pairs)

    def unexpected_callback(*_):
        pytest.fail("结构一致时不应调用语义对齐")

    same_pairs = main._build_pairs(str(source_path), str(source_path), alignment_callback=unexpected_callback)
    assert len(same_pairs) == 3

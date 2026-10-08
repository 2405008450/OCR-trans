"""按完整语义重建证件译文；正文使用原生段落，不受 OCR 小框限制。"""
from __future__ import annotations

import io
import json
import re
import time
from pathlib import Path

import fitz
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.core.config import settings
from app.service.doc_translate_service import SUPPORTED_LANGUAGES
from .budget import bounded_call
from .models import parse_json, validate_bbox
from .qa import render_docx
from .translation import call_translation, call_vision, validate_translations

ROLES = {"kicker", "title", "paragraph", "field", "note", "date", "footer"}


def validate_document(payload, page, index):
    if not isinstance(payload, dict) or not isinstance(payload.get("blocks"), list) or not payload["blocks"]:
        raise ValueError("可编辑排版响应缺少正文")
    if len(payload["blocks"]) > 150:
        raise ValueError("页面语义块过多")
    result = {"blocks": [], "images": [], "warnings": []}
    seen = set()
    for number, item in enumerate(payload["blocks"]):
        if not isinstance(item, dict) or item.get("role") not in ROLES:
            raise ValueError("可编辑排版包含未知段落类型")
        text, ids = item.get("text"), item.get("line_ids")
        if not isinstance(text, str) or not text.strip() or len(text) > 10000 or not isinstance(ids, list):
            raise ValueError("可编辑排版正文或文字行关联无效")
        for lid in ids:
            if type(lid) is not int or not 0 <= lid < len(page.lines):
                raise ValueError(f"可编辑排版包含未知文字行：{lid}；行号范围 0..{len(page.lines) - 1}")
            seen.add(lid)
        if len(ids) != len(set(ids)):
            raise ValueError("同一语义块包含重复文字行")
        result["blocks"].append({"id": f"p{index + 1}b{number + 1}", "role": item["role"],
                                 "source": text.strip(), "line_ids": ids, "translation": ""})
    # 装饰性底纹允许排除，但必须逐条说明；照片、印章覆盖的正文不能排除。
    excluded = payload.get("excluded_lines", [])
    if not isinstance(excluded, list):
        raise ValueError("排除文字行格式无效")
    for item in excluded:
        if not isinstance(item, dict) or item.get("reason") not in {"watermark", "decoration"}:
            raise ValueError("只能排除底纹或装饰文字")
        lid = item.get("line_id")
        if type(lid) is not int or not 0 <= lid < len(page.lines) or lid in seen:
            raise ValueError("排除文字行重复或无效")
        seen.add(lid)
        result["warnings"].append(f"第 {index + 1} 页装饰文字未纳入译文：{page.lines[lid].text}")
    if seen != set(range(len(page.lines))):
        raise ValueError(f"可编辑排版遗漏 OCR 文字行：{sorted(set(range(len(page.lines))) - seen)}")
    box = validate_bbox(payload.get("content_bbox"), page.width, page.height)
    result["landscape"] = box[2] - box[0] > box[3] - box[1]
    images = payload.get("images", [])
    if not isinstance(images, list) or len(images) > 4:
        raise ValueError("证件图片区域无效")
    for item in images:
        if not isinstance(item, dict) or item.get("kind") not in {"photo", "qr"}:
            raise ValueError("只允许保留原件照片或二维码，不生成印章与签名")
        result["images"].append({"kind": item["kind"], "bbox": validate_bbox(item.get("bbox"), page.width, page.height)})
    warnings = payload.get("warnings", [])
    if not isinstance(warnings, list) or any(not isinstance(w, str) or len(w) > 1000 for w in warnings):
        raise ValueError("辨识提醒格式无效")
    result["warnings"].extend(warnings[:20])
    return result


def extract_document(page, index, options):
    lines = [{"line_id": i, "text": line.text} for i, line in enumerate(page.lines)]
    prompt = (f"原图 {page.width}x{page.height} 像素。按阅读顺序完整抄录证件，重建自然段、字段与标题。"
              "OCR 只是参考，必须查看图片补回 OCR 漏识别的正文、日期、校名等。保持原语言，禁止提前翻译或编造。"
              "连接被分割的日期和句子；印章重叠不能导致整句正文被忽略。所有 OCR line_id 必须覆盖，行号从 0 开始，"
              "只能使用参考数据中的行号。若同一 OCR 行包含多个字段，允许分别关联到不同语义块，同一块内不能重复。"
              "补回的文字可用空 line_ids。照片仅提供紧贴照片的 bbox。印章与手写签名以 note 文字注明"
              "（例如：印章：原文机构名称；签名：可辨认姓名），不能复制整块带有中文正文的印章图片。"
              "无法辨认的内容保留疑问并加入 warnings，不能猜测。证书编号与网址逐字保留。"
              "content_bbox 是实际证件边界，忽略扫描纸张外围空白，用来判断横向或纵向。"
              '返回 {"content_bbox":[x1,y1,x2,y2],"blocks":[{"role":"kicker|title|paragraph|field|note|date|footer",'
              '"text":"原文完整语义段","line_ids":[0]}],"images":[{"kind":"photo|qr","bbox":[x1,y1,x2,y2]}],'
              '"excluded_lines":[{"line_id":1,"reason":"watermark|decoration"}],"warnings":[]}。'
              "禁止排除任何正文、编号、日期或校名。参考 OCR：" + json.dumps(lines, ensure_ascii=False))
    for attempt in range(2):
        payload = call_vision(page.image_path.read_bytes(), prompt, options["vision_model"], options["gemini_route"])
        (page.image_path.parent / f"editable-analysis-{index + 1}-{attempt}.json").write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        try:
            return validate_document(payload, page, index)
        except ValueError as exc:
            if attempt:
                raise
            prompt += ("\n上次响应未通过结构校验：" + str(exc) + "。仅纠正该问题，并重新核对原图，返回完整 JSON。上次响应：" +
                       json.dumps(payload, ensure_ascii=False))


def translate_documents(documents, options):
    language = SUPPORTED_LANGUAGES[options["target_lang"]]["english_name"]
    for document in documents:
        blocks = document["blocks"]
        for start in range(0, len(blocks), 40):
            batch = blocks[start:start + 40]
            prompt = (f"将以下完整证件段落从 {options['source_lang']} 翻译为 {language}。"
                      "段落之间保持姓名、机构、日期和术语一致。正常书面表达，不限制字符数，不缩写或省略事实。"
                      "日期可按目标语言转换写法，姓名用规范音译，编号和网址逐字保留。note 的印章和签名用文字注明。"
                      '返回 {"translations":[{"segment_id":"原 id","text":"完整译文"}]}。输入：' +
                      json.dumps([{"segment_id": b["id"], "role": b["role"], "text": b["source"]} for b in batch], ensure_ascii=False))
            values = validate_translations(parse_json(call_translation(prompt, options["translation_engine"], options["gemini_route"])), {b["id"] for b in batch})
            for block in batch:
                block["translation"] = values[block["id"]]


def _paragraph(container, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(0), Pt(6), 1.12
    fmt.widow_control = True
    run = paragraph.add_run(text)
    run.font.name, run.font.size = "Times New Roman", Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold, run.italic = bold, italic
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "SimSun")
    fonts.set(qn("w:cs"), "Arial")
    return paragraph


def _table(table, widths, border=False):
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "start", "end"):
                item = OxmlElement(f"w:{side}")
                item.set(qn("w:w"), "160" if border else "60")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            props.append(margins)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = Pt(1)
            cell.paragraphs[0].add_run().font.size = Pt(1)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:val"), "single" if border and side not in {"insideH", "insideV"} else "nil")
        item.set(qn("w:sz"), "8")
        item.set(qn("w:color"), "777777")
        borders.append(item)
    table._tbl.tblPr.append(borders)


def build_editable_docx(pages, documents, output_path, target_lang="en"):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    for index, (page, content) in enumerate(zip(pages, documents)):
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
        landscape = content["landscape"]
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Inches(11.693 if landscape else 8.268), Inches(8.268 if landscape else 11.693)
        section.top_margin = section.bottom_margin = Inches(0.4722)
        section.left_margin = section.right_margin = Inches(0.5903)
        width = (section.page_width - section.left_margin - section.right_margin) / 914400
        outer = document.add_table(rows=1, cols=1)
        _table(outer, [width], border=True)
        frame = outer.cell(0, 0)
        for block in content["blocks"]:
            if block["role"] in {"kicker", "title"}:
                p = _paragraph(frame, block["translation"], 25 if block["role"] == "title" else 16,
                               True, WD_ALIGN_PARAGRAPH.CENTER)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(12 if block["role"] == "title" else 4)
        body = frame
        if content["images"]:
            inner_width = width - 0.24
            table = frame.add_table(rows=1, cols=2)
            _table(table, [inner_width - 1.55, 1.55])
            body, pictures = table.cell(0, 0), table.cell(0, 1)
            for region in content["images"]:
                with Image.open(page.image_path) as source:
                    crop = source.crop(tuple(round(v) for v in region["bbox"]))
                    data = io.BytesIO()
                    crop.save(data, "PNG")
                    data.seek(0)
                    picture_width = min(1.25, 2.1 * crop.width / crop.height)
                    p = pictures.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.add_run().add_picture(data, width=Inches(picture_width))
        for block in content["blocks"]:
            role = block["role"]
            if role in {"kicker", "title", "footer"}:
                continue
            p = _paragraph(body, block["translation"], 11 if role == "note" else 12,
                           align=WD_ALIGN_PARAGRAPH.RIGHT if role == "date" else WD_ALIGN_PARAGRAPH.LEFT,
                           italic=role == "note")
            if target_lang in {"ar", "fa", "he", "ur"}:
                p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        for block in content["blocks"]:
            if block["role"] == "footer":
                p = _paragraph(frame, block["translation"], 9, align=WD_ALIGN_PARAGRAPH.RIGHT)
                p.paragraph_format.space_before = Pt(12)
        # 表格后的段落必须存在，但不占用额外一行正文高度。
        document.add_paragraph().paragraph_format.line_spacing = Pt(1)
    document.save(str(output_path))


def content_issues(pages, documents, options):
    issues = []
    for index, (page, content) in enumerate(zip(pages, documents)):
        source = "\n".join(line.text for line in page.lines) + "\n" + "\n".join(b["source"] for b in content["blocks"])
        translated = "\n".join(b["translation"] for b in content["blocks"])
        compact = re.sub(r"\s", "", translated)
        for token in set(re.findall(r"\d{6,}|https?://[^\s]+", source)):
            if token.rstrip("。.,；;") not in compact:
                issues.append({"page": index + 1, "kind": "identifier", "message": f"编号或网址需核对：{token}"})
        if options["source_lang"].startswith("zh") and options["target_lang"] not in {"zh", "zh-TW", "ja"}:
            for block in content["blocks"]:
                if re.search(r"[\u3400-\u9fff]", block["translation"]):
                    issues.append({"id": block["id"], "kind": "source_language", "message": "正文含未翻译中文"})
    return issues


def review_editable(pages, documents, output_path, debug_dir, options):
    issues = content_issues(pages, documents, options)
    report = {"status": "disabled", "output_mode": "editable", "issues": issues,
              "warnings": [w for d in documents for w in d["warnings"]], "visual_checked": False, "rounds": []}
    if not options["enable_qa"]:
        return report
    deadline = time.monotonic() + 180
    try:
        pdf = render_docx(output_path, Path(debug_dir) / "editable-qa", 45)
        with fitz.open(pdf) as rendered:
            if len(rendered) != len(pages):
                issues.append({"kind": "pagination", "message": f"原件 {len(pages)} 页，译文 {len(rendered)} 页，请复核分页"})
            for index, page in enumerate(pages):
                remaining = deadline - time.monotonic()
                if remaining < 5:
                    raise TimeoutError("译文质检总预算耗尽")
                if index >= len(rendered):
                    break
                with Image.open(page.image_path) as source:
                    original = source.convert("RGB")
                original.thumbnail((1300, 1500))
                pixmap = rendered[index].get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
                result = Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")
                result.thumbnail((1500, 1500))
                comparison = Image.new("RGB", (original.width + result.width, max(original.height, result.height)), "white")
                comparison.paste(original, (0, 0))
                comparison.paste(result, (original.width, 0))
                data = io.BytesIO()
                comparison.save(data, "PNG")
                comparison.save(Path(debug_dir) / f"editable-comparison-{index + 1}.png")
                prompt = ("左图原件，右图译文 Word 实际渲染。这是清晰可编辑的译文重排，允许改变文字位置和去除底纹；"
                          "印章与签名应以翻译后的文字注释代替。检查完整性与事实一致性，特别是姓名、日期、学校、专业、"
                          "编号、印章机构与签名，不要因为不同布局而报错。检查字号、溢出、遗漏和照片。"
                          '返回 {"issues":[{"kind":"事实或排版问题类型","message":"具体问题"}]}，无问题空数组。'
                          "以下为原文转录和译文：" + json.dumps(documents[index]["blocks"], ensure_ascii=False))
                review = bounded_call(call_vision, data.getvalue(), prompt,
                                      settings.LAYOUT_OVERLAY_QA_MODEL or options["vision_model"], options["gemini_route"],
                                      min(55, remaining - 2), budget_seconds=min(60, remaining))
                if not isinstance(review, dict) or not isinstance(review.get("issues"), list):
                    raise ValueError("译文质检响应无效")
                for issue in review["issues"][:30]:
                    if not isinstance(issue, dict) or not isinstance(issue.get("message"), str):
                        raise ValueError("译文质检问题格式无效")
                    issues.append({"page": index + 1, "kind": str(issue.get("kind", "review"))[:80], "message": issue["message"][:1500]})
        report["visual_checked"] = True
        report["status"] = "needs_review" if issues or report["warnings"] else "passed"
    except Exception as exc:
        report["status"] = "needs_review" if issues else "partial"
        report["warnings"].append(f"渲染或视觉质检未完成（{type(exc).__name__}），请复核 Word")
    return report
